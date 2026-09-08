#!/usr/bin/env python3
"""
グローバル株式 業種分類マップ → 株価パフォーマンス・ビューア用データ生成スクリプト

使い方:
  pip install yfinance pandas openpyxl
  python update_prices.py <業種分類マップ.xlsx>                # universe.js + prices.js を生成
  python update_prices.py <業種分類マップ.xlsx> --no-fundamentals  # P/E, P/B, 時価総額の取得をスキップ(高速)
  python update_prices.py <業種分類マップ.xlsx> --universe-only     # universe.js のみ再生成(価格取得なし)
  python update_prices.py <業種分類マップ.xlsx> --demo             # 乱数のデモデータで prices.js を生成(UI確認用)
  python update_prices.py <業種分類マップ.xlsx> --no-fundamentals --prev prices_prev.js
                                                                   # 株価だけ更新し、P/E等は前回分を引き継ぐ(定期実行用)

出力:
  universe.js  … 銘柄マスタ(セクター/業種グループ/国/時価総額/コメント)
  prices.js    … 期間別リターン(現地通貨・USD)、ベンチマークのリターン、P/E、P/B、時価総額

index.html と同じフォルダに置いてブラウザで開くだけで動作します(サーバー不要)。
"""
import argparse
import time
import datetime as dt
import json
import math
import re
import sys
from pathlib import Path

import openpyxl

CHUNK = 200          # 一度にダウンロードする銘柄数
PERIODS = {"1D": 1, "5D": 5, "1M": 21, "3M": 63, "6M": 126, "1Y": 252}  # 営業日数

# ---------------------------------------------------------------------------
# 国 → Yahoo Finance サフィックス / 通貨 / 各国代表指数
# ---------------------------------------------------------------------------
COUNTRY = {
    "米国":       dict(code="US", sfx="",    ccy="USD", idx="^GSPC"),
    "日本":       dict(code="JP", sfx=".T",  ccy="JPY", idx="^N225"),
    "台湾":       dict(code="TW", sfx=".TW", ccy="TWD", idx="^TWII"),
    "韓国":       dict(code="KR", sfx=".KS", ccy="KRW", idx="^KS11"),
    "中国":       dict(code="CN", sfx=".HK", ccy="HKD", idx="^HSI"),
    "インド":     dict(code="IN", sfx=".NS", ccy="INR", idx="^NSEI"),
    "マレーシア": dict(code="MY", sfx=".KL", ccy="MYR", idx="^KLSE"),
    "タイ":       dict(code="TH", sfx=".BK", ccy="THB", idx="^SET.BK"),
    "インドネシア": dict(code="ID", sfx=".JK", ccy="IDR", idx="^JKSE"),
    "フィリピン": dict(code="PH", sfx=".PS", ccy="PHP", idx="PSEI.PS"),
    "イギリス":   dict(code="GB", sfx=".L",  ccy="GBp", idx="^FTSE"),
    "ドイツ":     dict(code="DE", sfx=".DE", ccy="EUR", idx="^GDAXI"),
    "フランス":   dict(code="FR", sfx=".PA", ccy="EUR", idx="^FCHI"),
    "スイス":     dict(code="CH", sfx=".SW", ccy="CHF", idx="^SSMI"),
    "スウェーデン": dict(code="SE", sfx=".ST", ccy="SEK", idx="^OMX"),
    "イタリア":   dict(code="IT", sfx=".MI", ccy="EUR", idx="FTSEMIB.MI"),
    "オランダ":   dict(code="NL", sfx=".AS", ccy="EUR", idx="^AEX"),
    "スペイン":   dict(code="ES", sfx=".MC", ccy="EUR", idx="^IBEX"),
    "ノルウェー": dict(code="NO", sfx=".OL", ccy="NOK", idx="OBX.OL"),
    "デンマーク": dict(code="DK", sfx=".CO", ccy="DKK", idx="^OMXC25"),
    "フィンランド": dict(code="FI", sfx=".HE", ccy="EUR", idx="^OMXH25"),
    "ポーランド": dict(code="PL", sfx=".WA", ccy="PLN", idx="WIG20.WA"),
    "ベルギー":   dict(code="BE", sfx=".BR", ccy="EUR", idx="^BFX"),
    "オーストリア": dict(code="AT", sfx=".VI", ccy="EUR", idx="^ATX"),
    "アイルランド": dict(code="IE", sfx=".IR", ccy="EUR", idx="^ISEQ"),
    "ポルトガル": dict(code="PT", sfx=".LS", ccy="EUR", idx="PSI20.LS"),
}
INDEX_BENCH = {"S&P500": "^GSPC", "Nikkei225": "^N225", "STOXX600": "^STOXX", "MSCI Asia": "AAXJ"}
GLOBAL_BENCH = "ACWI"   # 対世界(USD)の基準
# 通貨 → USD換算用ペア(1通貨 = x USD)。Yahoo は "JPY=X" が USDJPY なので後で逆数を取る
FX_PAIRS = {"JPY", "TWD", "KRW", "HKD", "INR", "MYR", "THB", "IDR", "PHP", "GBP", "EUR",
            "CHF", "SEK", "NOK", "DKK", "PLN"}


def to_yahoo(ticker: str, country: str) -> str:
    """Bloomberg 風ティッカーを Yahoo Finance シンボルへ変換"""
    t = str(ticker).strip()
    c = COUNTRY[country]
    if country == "米国":
        return t.replace(".", "-")                       # BRK.B → BRK-B
    if country == "中国":
        if re.fullmatch(r"[A-Z]+", t):                  # 米国上場 ADR (PDD, TME, BZ ...)
            return t
        if re.fullmatch(r"\d{6}", t):                   # A株/B株
            return t + (".SS" if t[0] in "569" else ".SZ")
        return t.zfill(4) + ".HK"                       # 香港 (285 → 0285.HK)
    if country == "イギリス":
        t = t.rstrip(".").replace(".", "-")             # BT.A → BT-A, JD. → JD
    elif country in ("スウェーデン", "デンマーク", "フィンランド"):
        t = t.replace(" ", "-")                         # ERIC B → ERIC-B
    elif country == "タイ":
        t = t.replace(".R", "")                         # DELTA.R → DELTA
    return t + c["sfx"]


# ---------------------------------------------------------------------------
# Excel → universe
# ---------------------------------------------------------------------------
def read_universe(xlsx: Path):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["フラットDB"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    col = {h: i for i, h in enumerate(header)}
    out = []
    for r in rows[1:]:
        if not r[col["Ticker"]]:
            continue
        country = r[col["国"]]
        ticker = str(r[col["Ticker"]]).strip()
        cap = r[col["時価総額概算(B$)"]]
        wgt = r[col["指数ウェイト(%)"]]
        out.append({
            "id": f"{ticker}|{COUNTRY[country]['code']}",
            "t": ticker,
            "y": to_yahoo(ticker, country),
            "n": r[col["銘柄名"]],
            "c": COUNTRY[country]["code"],
            "cj": country,
            "ix": r[col["指数"]],
            "w": float(wgt) if isinstance(wgt, (int, float)) else None,
            "cap": float(cap) if isinstance(cap, (int, float)) else None,
            "s": r[col["セクター"]],
            "ss": r[col["サブセクター"]],
            "g": r[col["業種グループ"]],
            "d": r[col["コメント"]] or "",
        })
    # 注記シートの1行目(基準日など)を拾う
    note = ""
    try:
        note = str(wb["業種別ビュー"].cell(2, 1).value or "")
    except Exception:
        pass
    return out, note


def write_js(path: Path, var: str, obj):
    path.write_text(f"window.{var} = " + json.dumps(obj, ensure_ascii=False, separators=(",", ":")) + ";\n",
                    encoding="utf-8")
    print(f"  wrote {path}  ({path.stat().st_size/1024:.0f} KB)")


# ---------------------------------------------------------------------------
# 価格取得
# ---------------------------------------------------------------------------
def pct(series, n):
    """終値シリーズの直近値と n 営業日前の値からリターン(%)"""
    s = series.dropna()
    if len(s) <= n:
        return None
    a, b = float(s.iloc[-1]), float(s.iloc[-1 - n])
    if not b or math.isnan(a) or math.isnan(b):
        return None
    return round((a / b - 1) * 100, 2)


def load_prev(path):
    """前回の prices.js から P/E, P/B, 時価総額を引き継ぐ(株価だけ更新する高速モード用)"""
    if not path or not Path(path).exists():
        return {}
    try:
        t = Path(path).read_text(encoding="utf-8").strip()
        t = t[t.index("{"):].rstrip(";")
        return json.loads(t).get("prices", {})
    except Exception as e:
        print(f"  前回データを読めませんでした: {e}")
        return {}


def fetch_prices(universe, fundamentals=True, prev=None):
    import pandas as pd
    import yfinance as yf

    symbols = sorted({u["y"] for u in universe})
    bench = sorted(set(INDEX_BENCH.values()) | {GLOBAL_BENCH} | {c["idx"] for c in COUNTRY.values()})
    fx = sorted(f"{c}=X" for c in FX_PAIRS)

    print(f"  株価 {len(symbols)} 銘柄 / 指数 {len(bench)} / 為替 {len(fx)} をダウンロード中 ...")
    allsyms = symbols + bench + fx
    frames = []
    for i in range(0, len(allsyms), CHUNK):           # レート制限回避のため分割取得
        part = allsyms[i:i + CHUNK]
        for attempt in range(3):
            try:
                raw = yf.download(part, period="15mo", interval="1d", auto_adjust=True,
                                  group_by="column", threads=True, progress=False)
                break
            except Exception as e:
                print(f"    retry {attempt+1}: {e}")
                time.sleep(5 * (attempt + 1))
        else:
            continue
        if raw is None or raw.empty:
            continue
        c = raw["Close"] if isinstance(raw.columns, pd.MultiIndex) else raw[["Close"]].rename(columns={"Close": part[0]})
        frames.append(c)
        print(f"    {min(i+CHUNK, len(allsyms))}/{len(allsyms)}")
        time.sleep(1)
    close = pd.concat(frames, axis=1).sort_index().ffill(limit=5)

    def usd_factor(ccy):
        """1 現地通貨あたりの USD (時系列)。GBp はペンス建てなので 1/100"""
        if ccy == "USD":
            return None
        base = "GBP" if ccy == "GBp" else ccy
        col = f"{base}=X"
        if col not in close:
            return None
        f = 1.0 / close[col]
        return f / 100 if ccy == "GBp" else f

    def returns_for(series):
        return {p: pct(series, n) for p, n in PERIODS.items()}

    prices, missing = {}, []
    for u in universe:
        sym = u["y"]
        if sym not in close or close[sym].dropna().empty:
            missing.append(u["id"])
            continue
        s = close[sym]
        rec = {"loc": returns_for(s), "px": round(float(s.dropna().iloc[-1]), 2),
               "asof": str(s.dropna().index[-1].date())}
        f = usd_factor(COUNTRY[u["cj"]]["ccy"])
        if f is not None:
            rec["usd"] = returns_for((s * f.reindex(s.index).ffill()).dropna())
        else:
            rec["usd"] = rec["loc"]
        prices[u["id"]] = rec

    benches = {}
    for b in bench:
        if b in close and not close[b].dropna().empty:
            benches[b] = returns_for(close[b])

    if fundamentals:
        print("  P/E, P/B, 時価総額を取得中(銘柄ごとに問い合わせるため数十分かかります)...")
        fxlast = {c: (1.0 / float(close[f"{c}=X"].dropna().iloc[-1]) if f"{c}=X" in close else None)
                  for c in FX_PAIRS}
        for i, u in enumerate(universe):
            if u["id"] not in prices:
                continue
            info = {}
            for attempt in range(3):
                try:
                    info = yf.Ticker(u["y"]).info or {}
                    break
                except Exception:
                    time.sleep(2 * (attempt + 1))
            rec = prices[u["id"]]
            pe = info.get("trailingPE") or info.get("forwardPE")
            pb = info.get("priceToBook")
            mc = info.get("marketCap")
            rec["pe"] = round(float(pe), 1) if isinstance(pe, (int, float)) and pe > 0 else None
            rec["pb"] = round(float(pb), 2) if isinstance(pb, (int, float)) and pb > 0 else None
            if isinstance(mc, (int, float)) and mc > 0:
                ccy = info.get("currency") or COUNTRY[u["cj"]]["ccy"]
                if ccy == "USD":
                    k = 1.0
                elif ccy == "GBp":
                    k = (fxlast.get("GBP") or 0) / 100
                else:
                    k = fxlast.get(ccy) or 0
                if k:
                    rec["mcap"] = round(mc * k / 1e9, 1)
            if (i + 1) % 100 == 0:
                print(f"    {i+1}/{len(universe)}")
    elif prev:
        n = 0
        for uid, rec in prices.items():
            old = prev.get(uid) or {}
            for k in ("pe", "pb", "mcap"):
                if old.get(k) is not None:
                    rec[k] = old[k]
                    n += 1
        print(f"  前回データから P/E, P/B, 時価総額を引き継ぎました({n} 項目)")

    return {
        "generated": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "periods": list(PERIODS.keys()),
        "global_bench": GLOBAL_BENCH,
        "index_bench": INDEX_BENCH,
        "country_bench": {c["code"]: c["idx"] for c in COUNTRY.values()},
        "bench": benches,
        "prices": prices,
        "missing": missing,
        "demo": False,
    }


def demo_prices(universe):
    import random
    rnd = random.Random(7)
    scale = {"1D": 2, "5D": 4, "1M": 7, "3M": 12, "6M": 18, "1Y": 30}

    def rets():
        return {p: round(rnd.gauss(0, s), 2) for p, s in scale.items()}

    prices = {}
    for u in universe:
        r = rets()
        prices[u["id"]] = {"loc": r, "usd": {k: round(v + rnd.gauss(0, 1), 2) for k, v in r.items()},
                           "pe": round(abs(rnd.gauss(22, 12)) + 5, 1), "pb": round(abs(rnd.gauss(3, 2)) + 0.5, 2),
                           "px": round(abs(rnd.gauss(100, 50)) + 5, 2), "asof": str(dt.date.today())}
    bench = {b: {p: round(rnd.gauss(0, s / 2), 2) for p, s in scale.items()}
             for b in set(INDEX_BENCH.values()) | {GLOBAL_BENCH} | {c["idx"] for c in COUNTRY.values()}}
    return {"generated": dt.datetime.now().strftime("%Y-%m-%d %H:%M") + " (DEMO)", "periods": list(PERIODS),
            "global_bench": GLOBAL_BENCH, "index_bench": INDEX_BENCH,
            "country_bench": {c["code"]: c["idx"] for c in COUNTRY.values()},
            "bench": bench, "prices": prices, "missing": [], "demo": True}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", type=Path, help="業種分類マップの Excel ファイル")
    ap.add_argument("--out", type=Path, default=Path(__file__).parent, help="出力フォルダ(既定: このスクリプトの場所)")
    ap.add_argument("--no-fundamentals", action="store_true", help="P/E, P/B, 時価総額の取得をスキップ")
    ap.add_argument("--universe-only", action="store_true", help="universe.js のみ生成")
    ap.add_argument("--demo", action="store_true", help="乱数のデモデータで prices.js を生成")
    ap.add_argument("--prev", type=Path, default=None,
                    help="前回の prices.js。--no-fundamentals のとき P/E, P/B, 時価総額を引き継ぐ")
    a = ap.parse_args()

    universe, note = read_universe(a.xlsx)
    print(f"銘柄マスタ: {len(universe)} 銘柄")
    write_js(a.out / "universe.js", "UNIVERSE", {"note": note, "source": a.xlsx.name, "rows": universe})
    if a.universe_only:
        return
    if a.demo:
        data = demo_prices(universe)
    else:
        try:
            import yfinance  # noqa: F401
        except ImportError:
            sys.exit("yfinance が見つかりません:  pip install yfinance")
        data = fetch_prices(universe, fundamentals=not a.no_fundamentals, prev=load_prev(a.prev))
        print(f"取得: {len(data['prices'])} 銘柄 / 未取得: {len(data['missing'])} 銘柄")
        if data["missing"]:
            print("  未取得(シンボル対応を確認してください): " + ", ".join(data["missing"][:30])
                  + (" ..." if len(data["missing"]) > 30 else ""))
    write_js(a.out / "prices.js", "PRICES", data)


if __name__ == "__main__":
    main()
